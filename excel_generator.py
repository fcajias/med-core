import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io

def generar_excel_en_memoria(conn):
    """
    Genera en tiempo real el libro Excel maestro consultando los datos en vivo
    de la base de datos SQLite y retorna los bytes del archivo .xlsx.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    cur = conn.cursor()

    # Estilos
    header_fill_blue = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_fill_green = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    header_fill_purple = PatternFill(start_color="4C1D95", end_color="4C1D95", fill_type="solid")
    header_fill_amber = PatternFill(start_color="92400E", end_color="92400E", fill_type="solid")

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    strike_font = Font(name="Calibri", size=10, strike=True, color="9CA3AF")

    thin_border_side = Side(style='thin', color="D1D5DB")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    fill_green_badge = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    fill_yellow_badge = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    fill_red_badge = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fill_gray_badge = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")

    # -------------------------------------------------------------
    # 1. HOJA: INVENTARIO ACTUAL
    # -------------------------------------------------------------
    ws_inv = wb.create_sheet(title="INVENTARIO_ACTUAL")
    inv_headers = [
        "CÓDIGO", "MEDICAMENTO / INSUMO", "PRESENTACIÓN", "CONCENTRACIÓN", 
        "CATEGORÍA", "MARCAS COMERCIALES / SINÓNIMOS", "STOCK ACTUAL (BODEGA)", 
        "STOCK MÍNIMO", "ESTADO ALERTA"
    ]
    ws_inv.append(inv_headers)
    for col_idx in range(1, len(inv_headers) + 1):
        cell = ws_inv.cell(row=1, column=col_idx)
        cell.fill = header_fill_blue
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    cur.execute("""
        SELECT codigo, nombre, presentacion, concentracion, categoria, marcas_comerciales, stock_actual, stock_minimo
        FROM medicamentos
        WHERE activo = 1
        ORDER BY categoria DESC, nombre ASC
    """)

    for row_idx, r in enumerate(cur.fetchall(), start=2):
        cod, nom, pres, conc, cat, marcas, stock, smin = r
        if stock == 0:
            estado = "AGOTADO"
            badge_fill = fill_red_badge
        elif stock <= smin:
            estado = "STOCK BAJO"
            badge_fill = fill_yellow_badge
        else:
            estado = "DISPONIBLE"
            badge_fill = fill_green_badge
            
        ws_inv.append([cod, nom, pres, conc, cat, marcas, stock, smin, estado])
        for c_idx in range(1, 10):
            cell = ws_inv.cell(row=row_idx, column=c_idx)
            cell.font = cell_font
            cell.border = thin_border
            if c_idx in (1, 7, 8):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 9:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.fill = badge_fill
                cell.font = bold_font

    # -------------------------------------------------------------
    # 2. HOJA: PARTE DIARIO
    # -------------------------------------------------------------
    ws_pd = wb.create_sheet(title="PARTE_DIARIO")
    pd_headers = [
        "ID", "FECHA", "ESTADO", "CÉDULA", "PACIENTE", "EDAD", "CELULAR", 
        "PISO / ÁREA", "DIAGNÓSTICO", "MEDICAMENTOS ENTREGADOS", "TOTAL UNID", 
        "REGISTRADO POR", "MOTIVO ANULACIÓN"
    ]
    ws_pd.append(pd_headers)
    for col_idx in range(1, len(pd_headers) + 1):
        cell = ws_pd.cell(row=1, column=col_idx)
        cell.fill = header_fill_green
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    cur.execute("""
        SELECT a.id, a.fecha, COALESCE(a.estado, 'ACTIVA'), COALESCE(p.cedula, 'S/C'), 
               p.nombres || ' ' || p.apellidos, COALESCE(p.edad, ''), COALESCE(p.celular, ''), 
               COALESCE(p.piso_area, ''), a.diagnostico, COALESCE(a.usuario_registro, 'enfermeria'),
               COALESCE(a.motivo_anulacion, '')
        FROM atenciones a
        JOIN pacientes p ON a.paciente_id = p.id
        ORDER BY a.fecha ASC, a.id ASC
    """)
    atenciones = cur.fetchall()

    for row_idx, at in enumerate(atenciones, start=2):
        aid, f, est, ced, pac, edad, cel, piso, diag, u_reg, mot_anul = at
        cur.execute("""
            SELECT m.nombre || ' (' || m.presentacion || ') x' || d.cantidad, d.cantidad
            FROM despachos d
            JOIN medicamentos m ON d.medicamento_id = m.id
            WHERE d.atencion_id = ?
        """, (aid,))
        desp_list = cur.fetchall()
        med_text = ", ".join([d[0] for d in desp_list]) if desp_list else "Consulta / Procedimiento"
        tot_cant = sum([d[1] for d in desp_list])
        
        ws_pd.append([aid, f, est, ced, pac, edad, cel, piso, diag, med_text, tot_cant, u_reg, mot_anul])
        
        es_anulada = (est == "ANULADA")
        for c_idx in range(1, 14):
            cell = ws_pd.cell(row=row_idx, column=c_idx)
            cell.font = strike_font if es_anulada else cell_font
            cell.border = thin_border
            if c_idx in (1, 2, 3, 4, 6, 7, 8, 11):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if c_idx == 3:
                cell.fill = fill_red_badge if es_anulada else fill_green_badge
                cell.font = bold_font

    # -------------------------------------------------------------
    # 3. HOJA: KARDEX MOVIMIENTOS
    # -------------------------------------------------------------
    ws_kar = wb.create_sheet(title="KARDEX_MOVIMIENTOS")
    kar_headers = [
        "ID KARDEX", "FECHA", "CÓDIGO", "MEDICAMENTO / INSUMO", "TIPO MOVIMIENTO", 
        "DETALLE / CONCEPTO", "CANTIDAD", "STOCK ANTERIOR", "STOCK NUEVO", "REGISTRADO POR"
    ]
    ws_kar.append(kar_headers)
    for col_idx in range(1, len(kar_headers) + 1):
        cell = ws_kar.cell(row=1, column=col_idx)
        cell.fill = header_fill_purple
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    cur.execute("""
        SELECT k.id, k.fecha, m.codigo, m.nombre || ' (' || m.presentacion || ')', 
               k.tipo_movimiento, k.concepto, k.cantidad, k.stock_anterior, k.stock_nuevo,
               COALESCE(k.usuario_registro, 'sistema')
        FROM kardex k
        JOIN medicamentos m ON k.medicamento_id = m.id
        ORDER BY k.fecha ASC, k.id ASC
    """)

    for row_idx, kr in enumerate(cur.fetchall(), start=2):
        ws_kar.append(list(kr))
        t_mov = kr[4]
        for c_idx in range(1, 11):
            cell = ws_kar.cell(row=row_idx, column=c_idx)
            cell.font = cell_font
            cell.border = thin_border
            if c_idx in (1, 2, 3, 5, 7, 8, 9):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if c_idx == 5 and t_mov == "REVERSION_ANULACION":
                cell.fill = fill_yellow_badge
                cell.font = bold_font

    # -------------------------------------------------------------
    # 4. HOJA: ESTADISTICAS DE CONSUMO
    # -------------------------------------------------------------
    ws_st = wb.create_sheet(title="ESTADISTICAS_CONSUMO")
    ws_st.append(["TOP MEDICAMENTOS MÁS CONSUMIDOS (EXCLUYE CONSULTAS ANULADAS)"])
    ws_st.merge_cells("A1:D1")
    ws_st["A1"].font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    ws_st["A1"].fill = header_fill_amber
    ws_st["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws_st.append(["CÓDIGO", "MEDICAMENTO", "PRESENTACIÓN", "TOTAL UNIDADES CONSUMIDAS REALES"])
    for c_idx in range(1, 5):
        cell = ws_st.cell(row=2, column=c_idx)
        cell.font = header_font
        cell.fill = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Sumar solo atenciones activas
    cur.execute("""
        SELECT m.codigo, m.nombre, m.presentacion, sum(d.cantidad) as total_disp
        FROM despachos d
        JOIN medicamentos m ON d.medicamento_id = m.id
        JOIN atenciones a ON d.atencion_id = a.id
        WHERE a.estado != 'ANULADA'
        GROUP BY m.id
        ORDER BY total_disp DESC
        LIMIT 15
    """)

    for row_idx, r in enumerate(cur.fetchall(), start=3):
        ws_st.append(list(r))
        for c_idx in range(1, 5):
            cell = ws_st.cell(row=row_idx, column=c_idx)
            cell.font = cell_font
            cell.border = thin_border
            if c_idx in (1, 3, 4):
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Ajustar ancho de columnas
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 52)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()
